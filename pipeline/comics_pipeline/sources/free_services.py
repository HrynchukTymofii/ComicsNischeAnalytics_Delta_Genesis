"""Free-services workbook (data/raw/genesis/04_Безкоштовні сервіси/FREE_*.xlsx), collected 13.10.2026:
Similarweb (web traffic only, no apps), AppMagic demo (lower bounds), US store pages, Meta Ad Library,
Google Trends, WEBTOON SEC filings, deals, AlternativeTo and a list of data gaps.

The workbook is a human-readable report (titles, merged headers, notes under tables), so every table
is located by its header text, never by cell position; a missing header raises `LayoutError`.
The «Профіль …» tabs repeat the other tabs and are not read.
"""
from __future__ import annotations

import re
import warnings
from pathlib import Path

from openpyxl import load_workbook

from .common import AppMatcher, clean, num

SIMILARWEB = "Similarweb"
APPMAGIC = "AppMagic"


class LayoutError(ValueError):
    pass


def _rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def table(ws, first: str, second: str | None = None, stop=None) -> list[dict]:
    """Rows under the header row whose first cell is `first` (and second cell `second`, if given),
    keyed by header text, up to the first empty row or a row where `stop(row)` is true."""
    rows = _rows(ws)
    try:
        h = next(i for i, r in enumerate(rows) if clean(r[0]) == first and (second is None or clean(r[1]) == second))
    except StopIteration:
        raise LayoutError(f"sheet «{ws.title}»: header row starting «{first}» not found") from None
    header = [clean(c) for c in rows[h]]
    out = []
    for r in rows[h + 1:]:
        if not any(clean(c) for c in r) or (stop and stop(r)):
            break
        out.append({k: v for k, v in zip(header, r) if k})
    return out


def deal_subject(event) -> str:
    """'Kakao Entertainment → Tapas Media' -> 'Tapas Media' (the bought company); 'GlobalComix — раунд' -> 'GlobalComix'."""
    e = str(event or "")
    return e.split("→")[-1].strip() if "→" in e else e.split(" — ")[0].strip()


def _collected_on(wb) -> str | None:
    for r in _rows(wb["Читати першим"]):
        if clean(r[0]) == "Зібрано":
            m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", str(r[1]))
            if m:
                return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    return None


def _pairs(row: dict, prefix: str, value_prefix: str, n: int = 5) -> list[dict]:
    return [{"name": row[f"{prefix} {i}"], "share": num(row.get(f"{value_prefix} {i}"))}
            for i in range(1, n + 1) if clean(row.get(f"{prefix} {i}"))]


def read(path: Path, match: AppMatcher) -> dict:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl: unsupported conditional-formatting extensions
        wb = load_workbook(path, data_only=True)
    as_of = _collected_on(wb)
    sw_period = "2026-09"  # «вересень 2026» in every Similarweb tab
    out: dict = {"file": path.name, "as_of": as_of}

    # --- notes / methodology
    out["readme"] = [{"key": clean(r[0]), "text": clean(r[1])} for r in _rows(wb["Читати першим"])[1:] if clean(r[0]) and clean(r[1])]

    # --- brands (one row per player), pricing model, notes
    comp = table(wb["Конкуренти"], "Бренд", "Тип", stop=lambda r: clean(r[0]) == "Як читати")
    out["brands"] = [{
        "brand": r["Бренд"], "slug": match.resolve(name=r["Бренд"]), "type": clean(r.get("Тип")),
        "owner": clean(r.get("Власник")), "monetization": clean(r.get("Ціни і модель")),
        "ad_keyword": clean(r.get("Ключове слово")), "active_ads": num(r.get("Активних оголошень")),
        "advertisers": clean(r.get("Хто рекламується")), "trends_2026": num(r.get("Інтерес 2026")),
        "deals": clean(r.get("Угоди / фінанси")), "note": clean(r.get("Нотатка")),
    } for r in comp]
    out["how_to_read"] = [str(r[0]).lstrip("• ").strip() for r in _rows(wb["Конкуренти"]) if str(r[0] or "").startswith("•")]

    # --- Similarweb: traffic, countries, channels, social
    web = {}
    for r in table(wb["Similarweb трафік"], "Домен"):
        d = r["Домен"]
        web[d] = {"domain": d, "slug": match.resolve(name=d), "type": clean(r.get("Тип")), "source": SIMILARWEB,
                  "period": sw_period, "region": "world", "as_of": as_of,
                  "visits": num(r.get("Візити / міс")), "mom": num(r.get("Зміна до серпня")),
                  "mobile_share": num(r.get("Мобільний веб")), "desktop_share": num(r.get("Десктоп")),
                  "us_rank": num(r.get("Ранг у США")), "visit_duration": clean(r.get("Тривалість візиту")),
                  "pages_per_visit": num(r.get("Сторінок / візит")), "bounce_rate": num(r.get("Відмови")),
                  "x_webtoons": num(r.get("× від webtoons.com"))}
    for r in table(wb["Країни"], "Домен"):
        web.setdefault(r["Домен"], {"domain": r["Домен"], "slug": match.resolve(name=r["Домен"])})["countries"] = _pairs(r, "Країна", "Частка")
    channel_cols = ["Direct", "Org. Search", "Paid Search", "Referrals", "Display", "Org. Social", "Paid Social", "Gen AI", "Email", "Affiliates"]
    for r in table(wb["Канали"], "Домен"):
        web.setdefault(r["Домен"], {"domain": r["Домен"], "slug": match.resolve(name=r["Домен"])})["channels"] = {c: num(r.get(c)) for c in channel_cols}
    for r in table(wb["Соцмережі"], "Домен"):
        w = web.setdefault(r["Домен"], {"domain": r["Домен"], "slug": match.resolve(name=r["Домен"])})
        w["social"] = _pairs(r, "Мережа", "Частка")
        w["social_note"] = clean(r.get("Примітка"))
    tapas_details = table(wb["tapas.io деталі"], "Блок")
    if "tapas.io" in web:
        web["tapas.io"]["details"] = [{"block": r["Блок"], "metric": clean(r.get("Показник")), "value": r.get("Значення"),
                                       "comment": clean(r.get("Коментар"))} for r in tapas_details]
    out["web"] = list(web.values())

    # --- AppMagic (lower bounds), overlaps, similar apps
    am_rows = _rows(wb["AppMagic"])
    am = table(wb["AppMagic"], "Застосунок")
    out["appmagic"] = [{
        "source": APPMAGIC, "as_of": as_of, "period_days": 30, "period_note": clean(r.get("Джерело")),
        "app_name": r["Застосунок"], "slug": match.resolve(name=r["Застосунок"]), "publisher": clean(r.get("Видавець")),
        "revenue": num(r.get("Виручка 30 днів, $ (не менше)")), "downloads": num(r.get("Завантаження 30 днів (не менше)")),
        "bound": "lower", "store": "unknown", "region": "world",
        "revenue_by_country": clean(r.get("Виручка за країнами")), "downloads_by_country": clean(r.get("Завантаження за країнами")),
    } for r in am]
    overlaps, similar, current, similar_for = [], [], None, None
    for r in am_rows:
        a = clean(r[0])
        if not a:
            continue
        m = re.match(r"Конкуренти (.+?) \(overlap", a)
        if m:
            current = m.group(1)
        elif a.startswith("«Similar apps» для"):
            current = None
            similar_for = re.search(r"для (.+)$", a).group(1)
        elif current and isinstance(r[1], str) and r[1].endswith("%"):
            overlaps.append({"app": current, "app_slug": match.resolve(name=current), "competitor": a,
                             "competitor_slug": match.resolve(name=a), "overlap": num(r[1])})
        elif current is None and similar_for and "," in a and not clean(r[1]):
            names, _, comment = a.partition(" — ")
            similar.append({"app": similar_for, "app_slug": match.resolve(name=similar_for), "comment": comment or None,
                            "similar": [s.strip() for s in re.split(r",\s*(?![^()]*\))", names) if s.strip()]})
    out["appmagic_overlap"], out["appmagic_similar"] = overlaps, similar

    # --- US store pages
    out["stores_us"] = [{
        "app_name": r["Застосунок"], "slug": match.resolve(name=r["Застосунок"]), "as_of": as_of, "region": "US",
        "gp_rating": num(r.get("GP рейтинг")), "gp_reviews": num(r.get("GP відгуків")), "gp_installs": clean(r.get("GP інсталяцій")),
        "gp_note": clean(r.get("GP примітка")), "ios_rating": num(r.get("iOS рейтинг")), "ios_ratings": num(r.get("iOS оцінок")),
        "ios_chart": clean(r.get("iOS чарт")), "ios_iap": clean(r.get("iOS In-App Purchases (ціни)")),
    } for r in table(wb["Стори (US)"], "Застосунок")]

    # --- Meta Ad Library (US, active ads by keyword) + Manta hooks
    meta_rows = _rows(wb["Meta Ad Library"])
    out["meta_ad_library"] = [{
        "keyword": r["Ключове слово"], "active_ads": num(r.get("Активних оголошень (≈)")),
        "top_advertisers": clean(r.get("Хто рекламується (у топі)")), "region": "US", "as_of": as_of,
    } for r in table(wb["Meta Ad Library"], "Ключове слово", stop=lambda r: str(r[0] or "").startswith("Хуки"))]
    hooks_title = next((clean(r[0]) for r in meta_rows if str(r[0] or "").startswith("Хуки")), None)
    out["ad_hooks"] = {
        "title": hooks_title, "slug": match.resolve(name=re.search(r"Хуки (\S+)", hooks_title).group(1)) if hooks_title else None,
        "hooks": [str(r[0]).lstrip("• ").strip() for r in meta_rows if str(r[0] or "").startswith("•")],
        "angles": next((clean(r[0]) for r in meta_rows if str(r[0] or "").startswith("Кути")), None),
    }

    # --- Google Trends: several sets, each with its own header row «Рік | kw1 | kw2 …»
    sets, cur, set_title = [], None, None
    for r in _rows(wb["Google Trends"]):
        a = clean(r[0])
        if a == "Рік":
            cur = {"title": set_title, "keywords": [clean(c) for c in r[1:] if clean(c) and clean(c) != "Джерело"], "years": []}
            sets.append(cur)
        elif cur and a and re.fullmatch(r"\d{4}", a):
            cur["years"].append({"year": int(a), **{k: num(v) for k, v in zip(cur["keywords"], r[1:])}})
        elif a and a.startswith("Набір"):
            cur, set_title = None, a.rstrip(":")
        elif a and sets and not re.fullmatch(r"\d{4}", a) and not a.startswith("Зміна"):
            sets[-1]["note"] = a
    for s in sets:
        s["source"] = "Google Trends: США, Web Search, середнє за рік (2026 — до 11.10)"
    out["google_trends"] = sets

    # --- WEBTOON SEC, deals, AlternativeTo, gaps
    out["sec_webtoon"] = {"slug": match.resolve(name="WEBTOON"), "rows": [{
        "metric": r["Показник"], "values": {k: clean(r.get(k)) for k in ("FY2024", "FY2025", "Q2 2025", "Q2 2026")},
        "numbers": {k: num(r.get(k)) for k in ("FY2024", "FY2025", "Q2 2025", "Q2 2026")},
        "comment": clean(r.get("Коментар")), "source": clean(r.get("Джерело")),
    } for r in table(wb["SEC — WEBTOON"], "Показник")]}
    out["deals"] = [{
        "date": clean(r["Дата"]), "event": clean(r.get("Подія")), "type": clean(r.get("Тип")), "amount_usd": num(r.get("Сума, $")),
        "party": clean(r.get("Сторона / деталі")), "source": clean(r.get("Джерело")), "note": clean(r.get("Примітка")),
        "slugs": [s] if (s := match.by_name(deal_subject(r.get("Подія")))) else [],
    } for r in table(wb["Угоди і гроші"], "Дата")]
    out["alternativeto"] = [{
        "filter": clean(r.get("Фільтр")), "app": clean(r.get("Застосунок")), "likes": num(r.get("Лайки")),
        "price": clean(r.get("Ціна")), "country": clean(r.get("Країна")), "note": clean(r.get("Примітка")),
    } for r in table(wb["AlternativeTo"], "Фільтр")]
    out["alternativeto_for"] = match.resolve(name="Tapas")
    out["gaps"] = [{"service": clean(r.get("Сервіс")), "missing": clean(r.get("Чого бракує")),
                    "reason": clean(r.get("Причина")), "fix": clean(r.get("Як закрити"))} for r in table(wb["Прогалини"], "Сервіс")]
    return out
